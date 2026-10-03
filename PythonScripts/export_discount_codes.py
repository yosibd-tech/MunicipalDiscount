"""
Writes the current discount codes out to the two workbooks the importer reads.

    python export_discount_codes.py <codes.xlsx> <rates.xlsx>

    tbDiscountsCodes     -> codes.xlsx   one row per discount code
    tbDiscountCodeRates  -> rates.xlsx   one row per dated rate version

This is the round trip, and the way one municipality's code table is carried to
another: export here, hand over the two files, load them there with

    python import_discount_codes.py <codes.xlsx> <rates.xlsx> --replace

The column layout is identical to what import_discount_codes.py expects, so an
exported pair can be edited in Excel and loaded straight back. Neither workbook
holds a database id — a rate line finds its code through קוד עירוני + קוד חוקי,
which is what lets the target database assign its own ids.
"""

import os
import sys

import pyodbc
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# מחרוזת החיבור מגיעה מ-appsettings.json ומפוענחת ב-DPAPI —
# אין סיסמה בקוד. ראה db_config.py.
from db_config import CONN_STR

CONN_STR = CONN_STR
DEFAULT_CODES = os.path.join("Templates", "DiscountCodes.xlsx")
DEFAULT_RATES = os.path.join("Templates", "DiscountRates.xlsx")

# Header -> column, exactly as import_discount_codes.py reads them.
CODE_COLUMNS = [
    ("קוד עירוני",     "MunicipalDiscountCode", 12, None),
    ("קוד חוקי",       "LegallDiscountCode",    12, None),
    ("קטגוריה",        "Categories",            22, None),
    ("קוד קבוצה",      "GroupCode",             11, None),
    ("שם הנחה",        "DiscountName",          40, None),
    ("חובה",           "Mandatory",             10, None),
    ("חוק בתוקף",      "LawInForce",            16, None),
    ("מסמכים נדרשים",  "Documents",             32, None),
    ("הערה",           "Remark",                26, None),
    ("כללים",          "Rules",                 26, None),
    ("תנאים",          "Conditions",            26, None),
]

RATE_COLUMNS = [
    ("קוד עירוני",        "MunicipalDiscountCode",   12, None),
    ("קוד חוקי",          "LegallDiscountCode",      12, None),
    # Informational only - the import maps this header to nothing. Without it
    # the rates sheet is bare numbers, and the person adapting one
    # municipality's codes to another cannot tell which discount a row is.
    ("שם הנחה",           "DiscountName",            40, None),
    ("תאריך תחילת תוקף",  "ValidityDate",            18, "DD/MM/YYYY"),
    ("אחוז הנחה",         "DiscountPercent",         12, "0%"),
    ("אחוז הנחה קודם",    "PreviousDiscountPercent", 14, "0%"),
    ("מקסימום הנחה",      "MaximumDiscount",         13, "0%"),
    ("מקסימום שטח",       "MaximumArea",             13, "0.00"),
]

CODES_GUIDE = [
    ("קובץ קודי הנחה", True),
    ("שורה אחת לכל קוד הנחה. אין שנה ואין אחוזים — הם בקובץ התעריפים.", False),
    ("", False),
    ("חובה", True),
    ("קוד עירוני או קוד חוקי — לפחות אחד מהם בכל שורה.", False),
    ("צמד הקודים הוא המזהה של השורה, ולכן חייב להיות ייחודי.", False),
    ("", False),
    ("שים לב", True),
    ("קוד שאין לו אף שורה בקובץ התעריפים לא יימצא בשום חישוב.", False),
    ("כל עמודות הטקסט מוגבלות ל-255 תווים.", False),
    ("", False),
    ("העברה לרשות אחרת", True),
    ("שני הקבצים יחד הם התמונה המלאה. הטעינה ברשות היעד מוחקת את הקודים שם", False),
    ("ומחליפה אותם במה שיש כאן — אין מיזוג.", False),
]

RATES_GUIDE = [
    ("קובץ תעריפי הנחה", True),
    ("שורה אחת לכל גרסת תעריף.", False),
    ("קוד שלא השתנה — שורה אחת בלבד.", False),
    ("קוד שהעירייה שינתה במהלך השנה — שורה נוספת עם התאריך החדש.", False),
    ("", False),
    ("הקישור לקוד", True),
    ("קוד עירוני + קוד חוקי חייבים להתאים לשורה בקובץ קודי ההנחה.", False),
    ("עמודת שם ההנחה כאן היא לנוחות בלבד — הטעינה מתעלמת ממנה.", False),
    ("שינוי שם מתבצע בקובץ קודי ההנחה.", False),
    ("", False),
    ("אחוזים", True),
    ("העמודות מעוצבות כאחוז: הקלדת 40 תיתן 40%.", False),
    ("במסד הנתונים נשמר שבר (0.40).", False),
    ("", False),
    ("טעינה", True),
    ("python import_discount_codes.py <קודים> <תעריפים> --dry-run   בדיקה", False),
    ("python import_discount_codes.py <קודים> <תעריפים> --replace   טעינה", False),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)


def read_codes(cursor):
    cursor.execute("""
        SELECT ID, MunicipalDiscountCode, LegallDiscountCode, Categories, GroupCode,
               DiscountName, Mandatory, LawInForce, Documents, Remark, Rules, Conditions
        FROM   tbDiscountsCodes
        ORDER  BY ISNULL(MunicipalDiscountCode, 0), ISNULL(LegallDiscountCode, 0)""")

    codes = []
    by_id = {}
    for row in cursor.fetchall():
        entry = {
            "ID": row[0],
            "MunicipalDiscountCode": row[1],
            "LegallDiscountCode": row[2],
            "Categories": row[3],
            "GroupCode": row[4],
            "DiscountName": row[5],
            "Mandatory": row[6],
            "LawInForce": row[7],
            "Documents": row[8],
            "Remark": row[9],
            "Conditions": row[11],
            "Rules": row[10],
        }
        codes.append(entry)
        by_id[row[0]] = entry
    return codes, by_id


def read_rates(cursor, by_id):
    cursor.execute("""
        SELECT DiscountCodeId, ValidityDate, DiscountPercent,
               PreviousDiscountPercent, MaximumDiscount, MaximumArea
        FROM   tbDiscountCodeRates
        ORDER  BY DiscountCodeId, ValidityDate""")

    rates = []
    orphans = 0
    for row in cursor.fetchall():
        code = by_id.get(row[0])
        if code is None:          # cannot happen while the FK is in place
            orphans += 1
            continue
        rates.append({
            "MunicipalDiscountCode": code["MunicipalDiscountCode"],
            "LegallDiscountCode": code["LegallDiscountCode"],
            "DiscountName": code["DiscountName"],
            "ValidityDate": row[1],
            "DiscountPercent": float(row[2]) if row[2] is not None else None,
            "PreviousDiscountPercent": float(row[3]) if row[3] is not None else None,
            "MaximumDiscount": float(row[4]) if row[4] is not None else None,
            "MaximumArea": float(row[5]) if row[5] is not None else None,
        })

    if orphans:
        print(f"WARNING: {orphans} rate row(s) point at no code and were skipped")

    # Sort the way a person reads it: by code, newest version last.
    rates.sort(key=lambda r: (r["MunicipalDiscountCode"] or 0,
                              r["LegallDiscountCode"] or 0,
                              r["ValidityDate"]))
    return rates


def write_sheet(path, title, columns, rows, guide):
    book = Workbook()

    sheet = book.active
    sheet.title = title
    sheet.sheet_view.rightToLeft = True

    for index, (header, _, width, _) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "A2"

    for row_index, values in enumerate(rows, start=2):
        for col_index, (_, source, _, fmt) in enumerate(columns, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=values.get(source))
            if fmt:
                cell.number_format = fmt

    # Carry the formats into the empty rows below, so a code typed into a fresh
    # line behaves the same as one that came from the export.
    for blank_row in range(len(rows) + 2, len(rows) + 300):
        for col_index, (_, _, _, fmt) in enumerate(columns, start=1):
            if fmt:
                sheet.cell(row=blank_row, column=col_index).number_format = fmt

    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"

    notes = book.create_sheet("הוראות")
    notes.sheet_view.rightToLeft = True
    notes.column_dimensions["A"].width = 72

    for line_index, (text, is_title) in enumerate(guide, start=1):
        cell = notes.cell(row=line_index, column=1, value=text)
        if is_title:
            cell.font = Font(bold=True, size=12)

    book.save(path)


def main():
    codes_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_CODES
    rates_path = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_RATES

    for path in (codes_path, rates_path):
        folder = os.path.dirname(path)
        if folder:
            os.makedirs(folder, exist_ok=True)

    conn = pyodbc.connect(CONN_STR)
    try:
        cursor = conn.cursor()
        codes, by_id = read_codes(cursor)
        rates = read_rates(cursor, by_id)
    finally:
        conn.close()

    if not codes:
        print("ERROR: tbDiscountsCodes is empty — nothing to export.")
        sys.exit(1)

    without_rates = sum(1 for c in codes if not any(
        r["MunicipalDiscountCode"] == c["MunicipalDiscountCode"]
        and r["LegallDiscountCode"] == c["LegallDiscountCode"] for r in rates))
    if without_rates:
        print(f"WARNING: {without_rates} code(s) have no rate version and will not "
              f"be found by any calculation")

    write_sheet(codes_path, "קודי הנחה", CODE_COLUMNS, codes, CODES_GUIDE)
    write_sheet(rates_path, "תעריפי הנחה", RATE_COLUMNS, rates, RATES_GUIDE)

    print(f"Codes: {len(codes)} rows -> {codes_path}")
    print(f"Rates: {len(rates)} rows -> {rates_path}")


if __name__ == "__main__":
    main()
